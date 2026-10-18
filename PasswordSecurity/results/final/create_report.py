"""Convert the editable Markdown report into a polished Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT = Path(__file__).resolve().parents[2]
REPORT_SOURCE = PROJECT / "FINAL-REPORT.md"
BOARD_SOURCE = PROJECT / "DISPLAY-BOARD.md"
FINAL_DIRECTORY = Path(__file__).resolve().parent
NAVY = "14213D"
BLUE = "2563EB"
PALE = "E8F0FE"


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def add_table(document, markdown_rows):
    values = [[part.strip() for part in row.strip().strip("|").split("|")]
              for row in markdown_rows]
    table = document.add_table(rows=1, cols=len(values[0]))
    table.style = "Light Shading Accent 1"
    for index, value in enumerate(values[0]):
        cell = table.rows[0].cells[index]
        cell.text = value
        shade(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for values_row in values[2:]:
        cells = table.add_row().cells
        for index, value in enumerate(values_row):
            cells[index].text = value
    document.add_paragraph()


def add_image(document, source_line):
    match = re.match(r"!\[(.+?)\]\((.+?)\)", source_line)
    if not match:
        return
    caption, relative_path = match.groups()
    image_path = PROJECT / relative_path
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Inches(6.6))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.style = document.styles["Caption"]


def set_document_style(document):
    section = document.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.7)
    section.right_margin = Inches(.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [("Title", 28), ("Heading 1", 19),
                             ("Heading 2", 15), ("Heading 3", 12)]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if style_name == "Title" else BLUE)


def convert(source, output):
    document = Document()
    set_document_style(document)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    title_seen = False
    paragraph_lines = []

    def flush_paragraph():
        if paragraph_lines:
            document.add_paragraph(" ".join(line.strip() for line in paragraph_lines))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and "---" in lines[index + 1]:
            flush_paragraph()
            table_lines = [stripped, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, table_lines)
            continue
        if stripped.startswith("!["):
            flush_paragraph()
            add_image(document, stripped)
        elif stripped.startswith("# "):
            flush_paragraph()
            if title_seen:
                document.add_heading(stripped[2:], 0)
            else:
                heading = document.add_heading(stripped[2:], 0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_seen = True
        elif stripped.startswith("## "):
            flush_paragraph()
            document.add_heading(stripped[3:], 1)
        elif stripped.startswith("### "):
            flush_paragraph()
            document.add_heading(stripped[4:], 2)
        elif re.match(r"^\d+\. ", stripped):
            flush_paragraph()
            document.add_paragraph(re.sub(r"^\d+\. ", "", stripped), style="List Number")
        elif stripped.startswith("- "):
            flush_paragraph()
            document.add_paragraph(stripped[2:], style="List Bullet")
        elif stripped.startswith("**") and stripped.endswith("**"):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(stripped.strip("*")).bold = True
        else:
            paragraph_lines.append(stripped.replace("**", ""))
        index += 1
    flush_paragraph()

    # Add a small footer to every section.
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Password Security Science Fair Project • October 2026"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = RGBColor(91, 101, 119)

    document.save(output)
    # Reopen the package immediately so corruption is caught during generation.
    Document(output)
    print(output)


if __name__ == "__main__":
    convert(REPORT_SOURCE, FINAL_DIRECTORY / "Password-Security-Final-Report.docx")
    convert(BOARD_SOURCE, FINAL_DIRECTORY / "Password-Security-Display-Board.docx")
